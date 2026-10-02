"""Plain-text extraction from course documents (PDF, PPTX, DOCX, text formats).

Pure functions over bytes, with no Canvas or network access, so a caller can
run them in a worker thread. The PDF/PPTX/DOCX parsers come from the optional
``documents`` extra (``pip install 'canvas-mcp[documents]'``); when one is not
installed, ``MissingDependencyError`` names the package so the tool can answer
with an install hint instead of a traceback.

The bytes are third-party content. Office files are ZIP containers, so their
declared uncompressed size is checked before a parser inflates them, and
extraction can stop early once it has gathered a caller-chosen budget of text.
"""

import html
import io
import json
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import PurePosixPath

DOCUMENTS_EXTRA_HINT = (
    "Install the optional 'documents' extra on the machine running the MCP server: "
    "pip install 'canvas-mcp[documents]' (or: pip install pypdf python-pptx python-docx)."
)

#: Refuse an Office file whose members together inflate past this many bytes.
#: python-pptx/python-docx hold every part in memory, but non-XML parts
#: (images, media, embedded objects) are kept as raw bytes, roughly 1x their
#: size. Media in a deck barely compresses, so a legitimate file near the
#: 50 MB download cap can inflate a little past 50 MB; 2x leaves room for that.
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
#: Refuse when the XML parts together inflate past this many bytes. These are
#: the parts lxml parses in full before any text budget applies, at many times
#: their size in memory (a 190 MB document.xml peaked at ~3.4 GB), so this cap,
#: not the total, bounds parser memory. A long Word thesis or a 100-slide deck
#: stays well under it.
MAX_XML_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
#: Refuse a member that inflates more than this many times its compressed size
#: (once it is over ``_RATIO_CHECK_MIN_BYTES``). Real Office XML compresses
#: roughly 5-20x; a crafted part of repeated elements compresses hundreds of
#: times.
MAX_COMPRESSION_RATIO = 100
_RATIO_CHECK_MIN_BYTES = 1024 * 1024
_XML_SUFFIXES = (".xml", ".rels", ".vml")

KIND_PDF = "pdf"
KIND_PPTX = "pptx"
KIND_DOCX = "docx"
KIND_HTML = "html"
KIND_TEXT = "text"

#: Formats split into numbered pages or slides (``start_page`` applies).
PAGED_KINDS = frozenset({KIND_PDF, KIND_PPTX})

KIND_LABELS = {
    KIND_PDF: "PDF",
    KIND_PPTX: "PowerPoint (PPTX)",
    KIND_DOCX: "Word (DOCX)",
    KIND_HTML: "HTML",
    KIND_TEXT: "text",
}

SUPPORTED_FORMATS = (
    "PDF, PPTX, DOCX, and text files (plain text, Markdown, HTML, CSV, JSON, XML, "
    "LaTeX, notebooks, and source code)"
)

#: MIME types outside ``text/*`` whose content is text. Any ``+xml`` or
#: ``+json`` type counts too (see ``is_textual_type``).
TEXTUAL_TYPES = frozenset({
    "application/json",
    "application/csv",
    "application/xml",
    "application/javascript",
    "application/x-javascript",
    "application/ecmascript",
    "application/x-ipynb+json",
    "application/x-tex",
    "application/x-latex",
    "application/x-sh",
    "application/x-python",
    "application/x-python-code",
    "application/x-yaml",
    "application/yaml",
    "application/toml",
    "application/sql",
    "application/x-sql",
    "application/x-httpd-php",
    "application/x-perl",
    "application/x-ruby",
})

#: File extensions of text formats a course commonly hands out: notes, data,
#: markup, and starter code. Read as plain text.
TEXT_EXTENSIONS = frozenset({
    ".txt", ".text", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json",
    ".ipynb", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".log",
    ".tex", ".bib", ".sty", ".py", ".pyi", ".java", ".c", ".h", ".cpp", ".cc",
    ".cxx", ".hpp", ".hh", ".hxx", ".cs", ".js", ".mjs", ".cjs", ".ts", ".tsx",
    ".jsx", ".css", ".scss", ".go", ".rs", ".rb", ".php", ".pl", ".kt", ".kts",
    ".swift", ".scala", ".r", ".m", ".jl", ".lua", ".hs", ".ml", ".rkt",
    ".scm", ".lisp", ".clj", ".sql", ".sh", ".bash", ".zsh", ".ps1", ".bat",
    ".asm", ".s", ".v", ".sv", ".vhd", ".vhdl", ".mk", ".cmake", ".gradle",
    ".dockerfile", ".gitignore", ".srt", ".vtt",
})


def is_textual_type(mime: str) -> bool:
    """True when the MIME type names text content (a bare, lowercase token)."""
    return (
        mime.startswith("text/")
        or mime in TEXTUAL_TYPES
        or mime.endswith(("+xml", "+json"))
    )

_CONTENT_TYPES = {
    "application/pdf": KIND_PDF,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": KIND_PPTX,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": KIND_DOCX,
    "text/html": KIND_HTML,
    "application/xhtml+xml": KIND_HTML,
    "application/json": KIND_TEXT,
    "application/csv": KIND_TEXT,
}

_EXTENSIONS = {
    **dict.fromkeys(TEXT_EXTENSIONS, KIND_TEXT),
    ".pdf": KIND_PDF,
    ".pptx": KIND_PPTX,
    ".docx": KIND_DOCX,
    ".html": KIND_HTML,
    ".htm": KIND_HTML,
}


class DocumentTextError(Exception):
    """The document could not be turned into text; the message says why."""


class MissingDependencyError(DocumentTextError):
    """An optional parser package is not installed."""

    def __init__(self, package: str, kind: str) -> None:
        super().__init__(
            f"Reading {KIND_LABELS.get(kind, kind)} text needs the '{package}' package, "
            f"which is not installed. {DOCUMENTS_EXTRA_HINT}"
        )
        self.package = package


@dataclass
class TextSection:
    """One page, slide, or (for unpaged formats) the whole document."""

    number: int | None
    text: str


@dataclass
class ExtractedDocument:
    kind: str
    sections: list[TextSection]
    #: "page" or "slide" for paged formats, None otherwise.
    unit: str | None = None
    total_units: int | None = None
    #: True when extraction stopped early because the text budget was spent.
    stopped_early: bool = False
    notes: list[str] = field(default_factory=list)


def detect_kind(content_type: str | None, filename: str | None) -> str | None:
    """Pick an extractor from the Canvas content-type, then the file extension."""
    ctype = (content_type or "").split(";", 1)[0].strip().lower()
    if ctype in _CONTENT_TYPES:
        return _CONTENT_TYPES[ctype]
    if is_textual_type(ctype):
        return KIND_TEXT
    suffix = PurePosixPath((filename or "").lower()).suffix
    return _EXTENSIONS.get(suffix)


def _resolve_range(
    total: int, start: int | None, end: int | None, unit: str
) -> tuple[int, int]:
    """Clamp a 1-based inclusive range to ``1..total``."""
    first = start or 1
    last = min(end or total, total)
    if total == 0:
        return 1, 0
    if first > total:
        raise DocumentTextError(
            f"start_page {first} is past the last {unit} ({total} {unit}s in this file)."
        )
    return first, last


def _check_zip_container(data: bytes, kind: str) -> None:
    """Refuse an Office ZIP whose declared sizes would make parsing too costly.

    Uses the sizes in the ZIP central directory, so nothing is inflated here.
    A member that lies about its size fails the parser's CRC/size check
    instead of inflating past the declared value.
    """
    label = KIND_LABELS[kind]
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise DocumentTextError(
            f"The file is not a valid {label} document (not a ZIP container)."
        ) from exc

    inflated = 0
    xml_inflated = 0
    for info in members:
        inflated += info.file_size
        if info.filename.lower().endswith(_XML_SUFFIXES):
            xml_inflated += info.file_size
        if (
            info.file_size > _RATIO_CHECK_MIN_BYTES
            and info.file_size > MAX_COMPRESSION_RATIO * max(info.compress_size, 1)
        ):
            raise DocumentTextError(
                f"The {label} document has a part that expands more than "
                f"{MAX_COMPRESSION_RATIO}x when unpacked (a compression bomb pattern); "
                "refusing to parse it."
            )
    if inflated > MAX_UNCOMPRESSED_BYTES:
        raise DocumentTextError(
            f"The {label} document expands to more than "
            f"{MAX_UNCOMPRESSED_BYTES // (1024 * 1024)} MB when unpacked; refusing to parse it."
        )
    if xml_inflated > MAX_XML_UNCOMPRESSED_BYTES:
        raise DocumentTextError(
            f"The {label} document's XML expands to more than "
            f"{MAX_XML_UNCOMPRESSED_BYTES // (1024 * 1024)} MB when unpacked; "
            "refusing to parse it."
        )


def _extract_pdf(
    data: bytes, start: int | None, end: int | None, budget: int | None
) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise MissingDependencyError("pypdf", KIND_PDF) from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentTextError("The PDF is password-protected.")
        total = len(reader.pages)
    except DocumentTextError:
        raise
    except Exception as exc:
        raise DocumentTextError(f"Could not parse the PDF ({type(exc).__name__}).") from exc

    first, last = _resolve_range(total, start, end, "page")
    doc = ExtractedDocument(KIND_PDF, [], unit="page", total_units=total)
    used = 0
    empty = 0
    for number in range(first, last + 1):
        try:
            text = (reader.pages[number - 1].extract_text() or "").strip()
        except Exception as exc:
            text = f"[text extraction failed on this page: {type(exc).__name__}]"
        if not text:
            empty += 1
        doc.sections.append(TextSection(number, text))
        used += len(text)
        if budget is not None and used >= budget and number < last:
            doc.stopped_early = True
            break
    if doc.sections and empty == len(doc.sections):
        doc.notes.append(
            "No extractable text on these pages: the PDF is probably scanned images. "
            "This tool does not run OCR."
        )
    elif empty:
        doc.notes.append(f"{empty} page(s) had no extractable text (images or scans).")
    return doc


def _shape_text(shape: object) -> list[str]:
    """Text in a slide shape, descending into groups and tables."""
    lines: list[str] = []
    sub_shapes = getattr(shape, "shapes", None)
    if sub_shapes is not None:  # group shape
        for child in sub_shapes:
            lines.extend(_shape_text(child))
        return lines
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:  # type: ignore[attr-defined]
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        return lines
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()  # type: ignore[attr-defined]
        if text:
            lines.append(text)
    return lines


def _extract_pptx(
    data: bytes, start: int | None, end: int | None, budget: int | None
) -> ExtractedDocument:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise MissingDependencyError("python-pptx", KIND_PPTX) from exc

    _check_zip_container(data, KIND_PPTX)
    try:
        slides = list(Presentation(io.BytesIO(data)).slides)
    except Exception as exc:
        raise DocumentTextError(
            f"Could not parse the PowerPoint file ({type(exc).__name__})."
        ) from exc

    total = len(slides)
    first, last = _resolve_range(total, start, end, "slide")
    doc = ExtractedDocument(KIND_PPTX, [], unit="slide", total_units=total)
    used = 0
    for number in range(first, last + 1):
        slide = slides[number - 1]
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        body: list[str] = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            body.extend(_shape_text(shape))
        notes = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ""

        parts = []
        if title:
            parts.append(f"Title: {title}")
        parts.extend(body)
        if notes:
            parts.append(f"Speaker notes: {notes}")
        text = "\n".join(parts)
        doc.sections.append(TextSection(number, text))
        used += len(text)
        if budget is not None and used >= budget and number < last:
            doc.stopped_early = True
            break
    return doc


def _extract_docx(data: bytes, budget: int | None) -> ExtractedDocument:
    try:
        from docx import Document
        from docx.table import Table
    except ImportError as exc:
        raise MissingDependencyError("python-docx", KIND_DOCX) from exc

    _check_zip_container(data, KIND_DOCX)
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentTextError(f"Could not parse the Word file ({type(exc).__name__}).") from exc

    lines: list[str] = []
    used = 0
    stopped = False
    for block in document.iter_inner_content():
        added: list[str] = []
        if isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    added.append(" | ".join(cells))
        else:
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name or "") if block.style is not None else ""
            added.append(f"# {text}" if style.lower().startswith(("heading", "title")) else text)
        lines.extend(added)
        used += sum(len(line) + 1 for line in added)
        if budget is not None and used >= budget:
            stopped = True
            break
    return ExtractedDocument(
        KIND_DOCX, [TextSection(None, "\n".join(lines))], stopped_early=stopped
    )


class _HTMLText(HTMLParser):
    """Visible text of an HTML document, with line breaks at block elements."""

    _BLOCKS = frozenset({
        "p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
        "section", "article", "header", "footer", "table", "ul", "ol", "pre",
        "blockquote", "hr",
    })
    _SKIP = frozenset({"script", "style", "noscript", "template"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in self._BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)


def html_to_text(markup: str) -> str:
    parser = _HTMLText()
    parser.feed(markup)
    parser.close()
    text = html.unescape("".join(parser.parts))
    lines = [" ".join(line.split()) for line in text.splitlines()]
    out: list[str] = []
    for line in lines:
        if line or (out and out[-1]):
            out.append(line)
    return "\n".join(out).strip()


def decode_text(data: bytes) -> str:
    """Decode text bytes: UTF-8 (with or without BOM), else Windows-1252."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _extract_plain(data: bytes, kind: str, filename: str | None) -> ExtractedDocument:
    text = decode_text(data)
    if kind == KIND_HTML:
        text = html_to_text(text)
    elif (filename or "").lower().endswith(".json"):
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except ValueError:
            pass  # not valid JSON; return it as written
    return ExtractedDocument(kind, [TextSection(None, text.strip())])


def extract_text(
    data: bytes,
    kind: str,
    *,
    filename: str | None = None,
    start: int | None = None,
    end: int | None = None,
    budget: int | None = None,
) -> ExtractedDocument:
    """Extract text from ``data`` of the given ``kind`` (see ``detect_kind``).

    ``start``/``end`` are 1-based inclusive page (PDF) or slide (PPTX) numbers
    and are ignored, with a note, for unpaged formats. ``budget`` is an
    optional character count after which extraction stops early (setting
    ``stopped_early``); the default, None, extracts the whole document.
    """
    if kind == KIND_PDF:
        return _extract_pdf(data, start, end, budget)
    if kind == KIND_PPTX:
        return _extract_pptx(data, start, end, budget)
    if kind == KIND_DOCX:
        doc = _extract_docx(data, budget)
    elif kind in (KIND_TEXT, KIND_HTML):
        doc = _extract_plain(data, kind, filename)
    else:
        raise DocumentTextError(f"Unsupported document kind: {kind}")
    if start is not None or end is not None:
        doc.notes.append(
            f"{KIND_LABELS[kind]} files have no pages or slides, so the page range was ignored."
        )
    return doc
