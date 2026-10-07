"""Shared pytest fixtures for Canvas MCP tests."""

import io
import os
from unittest.mock import AsyncMock, patch

import pytest

# canvas_mcp.core.config calls load_dotenv() at import time, and python-dotenv
# walks up from the working directory, so a developer's real .env (or one in a
# parent checkout of a git worktree) would otherwise change test outcomes
# (write-tool allowlists, policy defaults, role). Must run before the first
# canvas_mcp import below.
os.environ["PYTHON_DOTENV_DISABLED"] = "1"

from canvas_mcp.core.config import reset_config  # noqa: E402


@pytest.fixture(autouse=True)
def reset_config_between_tests(monkeypatch):
    """Discard the cached config singleton before and after each test.

    Without this, the first test to call get_config() freezes the singleton
    from its environment; later tests that patch env vars (e.g. anonymization
    toggles) would silently read stale config.

    Also pins ACCESSIBILITY_CHECKERS to its default so a developer's .env or
    shell (e.g. ``ACCESSIBILITY_CHECKERS=none``, a supported setting since
    #325) cannot shrink the registry under the suite. Tests of the gate itself
    override this through their own fixture.
    """
    monkeypatch.setenv("ACCESSIBILITY_CHECKERS", "ufixit")
    reset_config()
    yield
    reset_config()


@pytest.fixture(autouse=True)
def isolated_course_cache(monkeypatch):
    """Start every test with an empty course cache and restore it afterwards.

    The course cache is per principal and process-global, so a test that
    refreshes it would otherwise leak its synthetic courses into later tests
    and change how ``resolve_numeric_course_id`` answers there. Resetting also
    drops the refresh-on-miss rate limit and any shared in-flight refresh, so
    one test's refresh never suppresses the next test's.
    """
    from canvas_mcp.core import cache

    cache.reset_course_cache()
    yield cache
    cache.reset_course_cache()


@pytest.fixture
def mock_canvas_request():
    """Mock Canvas API request function."""
    with patch('canvas_mcp.core.client.make_canvas_request') as mock:
        mock.return_value = AsyncMock()
        yield mock


@pytest.fixture
def mock_fetch_paginated():
    """Mock paginated fetch function."""
    with patch('canvas_mcp.core.client.fetch_all_paginated_results') as mock:
        mock.return_value = AsyncMock()
        yield mock


@pytest.fixture
def mock_course_id_resolver():
    """Mock course ID resolver."""
    with patch('canvas_mcp.core.cache.get_course_id') as mock:
        # Default to returning the input as-is (assuming it's already an ID)
        async def resolve_id(identifier):
            return str(identifier) if isinstance(identifier, int) else identifier
        mock.side_effect = resolve_id
        yield mock


@pytest.fixture
def mock_course_code_resolver():
    """Mock course code resolver."""
    with patch('canvas_mcp.core.cache.get_course_code') as mock:
        async def resolve_code(course_id):
            return f"course_{course_id}"
        mock.side_effect = resolve_code
        yield mock


@pytest.fixture
def sample_course_data():
    """Sample course data for testing."""
    return {
        "id": 12345,
        "name": "Introduction to Computer Science",
        "course_code": "CS101_2024",
        "start_at": "2024-01-15T08:00:00Z",
        "end_at": "2024-05-15T17:00:00Z",
        "time_zone": "America/Chicago",
        "default_view": "modules",
        "is_public": False,
        "blueprint": False
    }


@pytest.fixture
def sample_assignment_data():
    """Sample assignment data for testing."""
    return {
        "id": 67890,
        "name": "Python Programming Project",
        "description": "<p>Build a Python application</p>",
        "due_at": "2024-02-15T23:59:00Z",
        "points_possible": 100,
        "submission_types": ["online_upload", "online_text_entry"],
        "published": True,
        "locked_for_user": False
    }


@pytest.fixture
def sample_submission_data():
    """Sample submission data for testing."""
    return {
        "id": 111,
        "user_id": 1001,
        "submitted_at": "2024-02-14T18:30:00Z",
        "score": 85,
        "grade": "85",
        "workflow_state": "graded",
        "late": False,
        "missing": False,
        "excused": False
    }


@pytest.fixture
def sample_page_data():
    """Sample page data for testing."""
    return {
        "page_id": 222,
        "url": "module-1-overview",
        "title": "Module 1: Overview",
        "body": "<h1>Welcome to Module 1</h1><p>Content here</p>",
        "published": True,
        "front_page": False,
        "updated_at": "2024-01-20T10:00:00Z"
    }


@pytest.fixture
def sample_rubric_data():
    """Sample rubric data for testing."""
    return {
        "id": 333,
        "title": "Programming Assignment Rubric",
        "context_id": 12345,
        "context_type": "Course",
        "points_possible": 100,
        "criteria": [
            {
                "id": "crit1",
                "description": "Code Quality",
                "points": 40,
                "ratings": [
                    {"id": "r1", "description": "Excellent", "points": 40},
                    {"id": "r2", "description": "Good", "points": 30},
                    {"id": "r3", "description": "Fair", "points": 20},
                    {"id": "r4", "description": "Poor", "points": 0}
                ]
            },
            {
                "id": "crit2",
                "description": "Documentation",
                "points": 30,
                "ratings": [
                    {"id": "r5", "description": "Excellent", "points": 30},
                    {"id": "r6", "description": "Good", "points": 20},
                    {"id": "r7", "description": "Fair", "points": 10},
                    {"id": "r8", "description": "Poor", "points": 0}
                ]
            }
        ]
    }


@pytest.fixture
def sample_discussion_topic_data():
    """Sample discussion topic data for testing."""
    return {
        "id": 444,
        "title": "Week 1 Discussion",
        "message": "Discuss this week's topics",
        "posted_at": "2024-01-15T09:00:00Z",
        "published": True,
        "discussion_type": "threaded",
        "user_can_see_posts": True
    }


@pytest.fixture
def sample_announcement_data():
    """Sample announcement data for testing."""
    return {
        "id": 555,
        "title": "Important: Exam Schedule",
        "message": "<p>The midterm exam will be on March 1st</p>",
        "posted_at": "2024-02-01T12:00:00Z",
        "author": {"id": 2000, "display_name": "Prof. Smith"}
    }


# --- Course document builders (read_course_file_text tests) -----------------
# Built in-process so expected text is known independently of the extractor.
# python-pptx / python-docx come from the optional "documents" extra; tests
# that use those builders skip when it is not installed.


def _make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF whose page i shows the literal text pages[i].

    An empty string makes a page with no text (what a scanned page looks like
    to a text extractor).
    """
    objects: list[bytes] = []
    count = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(count))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                "/Resources << /Font << /F1 3 0 R >> >> "
                f"/Contents {5 + 2 * i} 0 R >>"
            ).encode()
        )
        objects.append(
            b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1) + b"0000000000 65535 f \n"
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def _make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    first.shapes.title.text = "Intro to Graphs"
    first.placeholders[1].text = "Vertices and edges"
    table = first.shapes.add_table(
        2, 2, Inches(1), Inches(4), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "BFS"
    table.cell(0, 1).text = "O(V+E)"
    first.notes_slide.notes_text_frame.text = "Mention the midterm"
    second = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    second.shapes.title.text = "Shortest Paths"
    third = prs.slides.add_slide(prs.slide_layouts[5])
    third.shapes.title.text = "Dijkstra"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Course Policies", 1)
    document.add_paragraph("Late work loses 10% per day.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Week 1"
    table.cell(0, 1).text = "Introduction"
    document.add_paragraph("Office hours are on Fridays.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()




@pytest.fixture
def make_pdf():
    """Builder: list of page strings -> PDF bytes ("" = a page with no text)."""
    return _make_pdf


@pytest.fixture
def make_pptx():
    """Builder: a 3-slide deck (title/body/table/notes, then two title slides)."""
    pytest.importorskip("pptx")
    return _make_pptx


@pytest.fixture
def make_docx():
    """Builder: heading, paragraph, table, paragraph."""
    pytest.importorskip("docx")
    return _make_docx
